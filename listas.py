from openpyxl import Workbook
from openpyxl.styles import Font
import os
meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

class GestorListas:
    @staticmethod
    def crearLista1(libro):
        """LISTA GNSS - RVSM"""
        hojaLista1 = libro.create_sheet('GNSS - RVSM')
        hojaLista1['A1'] = 'TYP'
        hojaLista1['B1'] = 'GNSS'
        hojaLista1['C1'] = 'RVSM'
        typ = ["A20N", "A319", "A320", "A332", "AC90", "B190", "B350", "B738", "B462", "B733",
               "B734", "B735", "B737", "B762", "B763", "B788", "B789", "BE20", "BE25", "BE40",
               "BE90", "BE9L", "C130", "C150", "C152", "C172", "C206", "C208", "C210", "C402",
               "C550", "C560", "C750", "CRJ2", "DA40", "DA42", "DC10", "DH8B", "E145", "E190",
               "F27", "F50", "F900", "FA50", "FA7X", "G150", "GALX", "GLEX", "GLF3", "GLF4",
               "GLF5", "K8W", "LJ25", "LJ45", "LJ60", "P56C", "PA23", "PA31", "PA32", "PA34",
               "PA46", "PC12", "PC7", "RJ1H", "RJ70", "RJ85", "RV4", "SR20", "T210", "A339",
               "A359", "ASTR", "BE200", "BE35", "C295", "C340", "C650", "C680", "CL60", "DH8",
               "E55P", "EC45", "FA8X", "GA8C", "TBM8", "PA27", "AC112", "C25B", "C441", "LJ35",
               "SW4", "JS32"]
        gnss = ["G", "G", "G", "G", "", "G", "G", "G", "G", "G",
                "G", "G", "G", "G", "G", "G", "G", "G", "", "G",
                "", "", "G", "", "", "G", "G", "G", "", "",
                "G", "G", "G", "G", "G", "G", "G", "G", "G", "G",
                "", "G", "G", "G", "G", "G", "G", "G", "G", "G",
                "G", "", "", "G", "G", "", "", "", "", "",
                "", "G", "", "G", "G", "G", "", "G", "", "G",
                "G", "G", "G", "", "G", "", "G", "G", "G", "G",
                "G", "G", "G", "", "G", "", "", "G", "G", "G",
                "G", "G"]
        rvsm = ["W", "W", "W", "W", "", "W", "W", "W", "W", "W",
                "W", "W", "W", "W", "W", "W", "W", "W", "", "W",
                "", "", "W", "", "", "", "", "", "", "",
                "W", "W", "W", "W", "W", "", "", "W", "W", "W",
                "", "W", "W", "W", "W", "W", "W", "W", "W", "W",
                "W", "", "", "W", "W", "", "", "", "", "",
                "", "W", "", "W", "W", "W", "", "", "", "W",
                "W", "W", "W", "", "W", "", "W", "W", "W", "W",
                "W", "", "W", "", "W", "", "", "W", "", "",
                "", ""]
        for i in range(len(typ)):
            hojaLista1[f'A{i+2}'] = typ[i]
            hojaLista1[f'B{i+2}'] = gnss[i]
            hojaLista1[f'C{i+2}'] = rvsm[i]

    @staticmethod
    def crearLista2(libro):
        """LISTA AD"""
        hojaLista2 = libro.create_sheet('LISTA AD')
        hojaLista2['A1'] = 'AD'
        ad = ["SLAL", "SLCO", "SLLP", "SLOR", "SLRQ", "SLTJ", "SLTR", "SLVR", "SLCB",
              "SLET", "LEMD", "SLTL", "SLRI", "SBSJ", "SAEZ", "SPJC", "SKLT", "SLHI", "SGAS",
              "SLGM", "SLUY", "SPQU", "MPMG", "SACO", "SANT", "SCEL", "MROC", "SLSM", "SBGR",
              "SUMU", "SBGO", "SLSU", "SLSA", "SBPA", "LFBO", "ZZZZ", "SCIE", "SBGL", "SLYA",
              "SBFZ", "SLSR", "SLSB", "SVMI", "SBKP", "SLVM", "LCL", "KMIA", "SBEG", "SKCL",
              "SPZO", "SBCG", "SBBR", "SVPA", "SLPO"]
        for i in range(len(ad)):
            hojaLista2[f'A{i+2}'] = ad[i]

    @staticmethod
    def crearLista3(libro):
        """LISTA REG"""
        hojaLista3 = libro.create_sheet('LISTA REG')
        hojaLista3['A1'] = 'REG'
        hojaLista3['B1'] = 'TYP'
        reg = ["CP3199", "CP2920", "CP2921", "CP3018", "CP2925", "CP2924", "EB003", "CP2926", "CP3138", "CP3019",
               "CP3145", "FAB61", "CP3196", "CP2889", "CP3151", "CP3194", "CP2790", "CP3206", "CP2791", "CP3108",
               "ECNBN", "CP2922", "FAB050", "CP3142", "CP2709", "CP3017", "FAB046", "CP3118", "YV1118", "N381VP",
               "CP2286", "CP2645", "CP3204", "CP2819", "CP2881", "FAB90", "HK5255", "CP2828", "CP2596", "CP2600",
               "N562PA", "CP3086", "LVCBK", "FAB661", "FAB660", "FAB409", "CP2923", "CP3144", "CP3121", "HCCUH",
               "FAB002", "XBVHL", "FAB453", "FAB470", "CP3171", "N488RJ", "CP2667", "CP3062", "CP2380", "ECMAJ",
               "YV2716", "CP3120", "LVFQF", "T99", "FAB004", "CP3193", "FAB001", "EB103", "CP3001", "CP1583",
               "CP2442", "CP3077", "CP3135", "CP3205", "CP3104", "CP2672", "FAB81", "LVGWV", "CP3208", "LVFVZ",
               "CP2852", "CP3106", "CP3000", "CP3143", "FAB021", "FAB047", "CP3101", "N225AX", "FAB664", "FAB662",
               "CP2646", "T7ESPRT", "FAB048", "CP2737", "CP2474", "CP2814", "LVFUF", "FAB362", "FAB008", "XAVBC",
               "1402", "CP3210", "CP3209", "CP2467", "CP2539", "CP3033", "FAB011", "CP2844", "LVWXD", "FAB106",
               "FAB018", "CP3158", "CP3212", "CP2768", "CP2723", "CP2265", "CPX3161", "CP2883", "FAB412", "CPX2403",
               "CP3214", "CP3215", "PB004", "LVGQR", "FAB86", "PB003"]
        typ = ["B738", "B733", "B733", "B737", "B738", "B737", "BE90", "B738", "B738", "B733",
               "E190", "C130", "B738", "RJ85", "B738", "RJ1H", "C550", "B738", "DC10", "B190",
               "A332", "B737", "BE90", "E190", "PA31", "B763", "JS32", "C208", "LJ45", "C553",
               "C206", "C172", "B738", "C172", "B763", "F27", "LJ45", "C172", "C172", "B350",
               "C525", "B763", "C560", "K8W", "K8W", "C206", "B737", "PA34", "B190", "B735",
               "FA50", "H25B", "PC7", "PC7", "E190", "H25B", "C152", "C172", "PA27", "A332",
               "LJ45", "C182", "C340", "B737", "EC45", "A60", "F900", "Z9EH", "C172", "C172",
               "C206", "B733", "E190", "PC12", "C206", "PA32", "C130", "BE20", "A332", "LJ60",
               "CRJ2", "RJ70", "C172", "EC45", "C402", "BE25", "C206", "B762", "K8W", "K8W",
               "C152", "F900", "BE35", "C172", "C172", "RJ85", "LJ60", "C210", "LJ25", "E145",
               "GL5T", "CRJ2", "A332", "AC90", "C172", "P68C", "BE20", "C206", "C551", "B462",
               "BE20", "C402", "PC12", "P46T", "C172", "C152", "RV4", "C172", "C210", "P56C",
               "A332", "B738", "C206", "LJ60", "C212", "BE55"]
        for i in range(len(reg)):
            hojaLista3[f'A{i+2}'] = reg[i]
            hojaLista3[f'B{i+2}'] = typ[i]

    @staticmethod
    def crearLista4(libro):
        """LISTA CAT - WTC"""
        hojaLista4 = libro.create_sheet('LISTA CAT - WTC')
        hojaLista4['A1'] = 'TYP'
        hojaLista4['B1'] = 'CAT'
        hojaLista4['C1'] = 'WTC'
        typ = ["A20N", "A319", "A320", "A332", "AC90", "B190", "B350", "B738", "B462", "B733",
               "B734", "B735", "B737", "B738", "B762", "B763", "B788", "B789", "BE20", "BE25",
               "BE40", "BE90", "BE9L", "C130", "C150", "C152", "C172", "C206", "C206", "C208",
               "C210", "C402", "C550", "C560", "C750", "CRJ2", "DA40", "DA42", "DC10", "DH8B",
               "E145", "E190", "F27", "F50", "F900", "FA50", "FA7X", "G150", "GALX", "GLEX",
               "GLF3", "GLF4", "GLF5", "K8W", "LJ25", "LJ45", "LJ60", "P56C", "PA23", "PA31",
               "PA32", "PA34", "PA46", "PC12", "PC7", "RJ1H", "RJ70", "RJ85", "RV4", "SR20",
               "T210", "A339", "A359", "ASTR", "BE200", "BE35", "C295", "C340", "C650", "C680",
               "CL60", "DH8", "E55P", "EC45", "FA8X", "GA8C", "TBM8", "PA27", "AC112", "C25B",
               "C441", "LJ35", "SW4", "JS32"]
        cat = ["C", "C", "C", "C", "B", "B", "B", "D", "C", "C",
               "C", "C", "C", "D", "C", "C", "C", "C", "B", "B",
               "B", "A", "A", "C", "A", "A", "A", "A", "A", "A",
               "A", "B", "B", "B", "C", "C", "A", "A", "D", "B",
               "C", "C", "B", "B", "B", "C", "B", "C", "C", "C",
               "C", "C", "C", "C", "C", "C", "C", "A", "A", "A",
               "A", "A", "A", "A", "A", "C", "C", "C", "A", "A",
               "B", "D", "D", "B", "A", "A", "C", "B", "B", "B",
               "C", "C", "B", "A", "C", "A", "A", "B", "B", "B",
               "B", "C", "B", "B"]
        wtc = ["M", "M", "M", "H", "L", "M", "L", "M", "M", "M",
               "M", "M", "M", "M", "H", "H", "H", "H", "L", "L",
               "M", "L", "L", "M", "L", "L", "L", "L", "L", "L",
               "L", "L", "L", "M", "M", "M", "L", "L", "H", "M",
               "M", "M", "M", "M", "M", "M", "M", "M", "M", "M",
               "M", "M", "M", "L", "L", "M", "M", "L", "L", "L",
               "L", "L", "L", "L", "L", "M", "M", "M", "L", "L",
               "L", "H", "H", "M", "", "L", "M", "L", "M", "M",
               "M", "M", "M", "L", "M", "L", "L", "L", "L", "L",
               "L", "M", "L", "M"]
        for i in range(len(typ)):
            hojaLista4[f'A{i+2}'] = typ[i]
            hojaLista4[f'B{i+2}'] = cat[i]
            hojaLista4[f'C{i+2}'] = wtc[i]

def crearFechaArchivo(indiceMes, year):
    """Crea todas las listas en un nuevo archivo Excel"""
    libro = Workbook()
    
    # Crear todas las listas
    GestorListas.crearLista1(libro)
    GestorListas.crearLista2(libro)
    GestorListas.crearLista3(libro)
    GestorListas.crearLista4(libro)
    
    # Eliminar la hoja por defecto si existe
    if 'Sheet' in libro.sheetnames:
        del libro['Sheet']
    
    año = year
    mes = meses[indiceMes]
    ruta_segura = os.path.join(os.path.expanduser('~'), 'Desktop', f'{mes}_{año}.xlsx')
    libro.save(ruta_segura)
    return ruta_segura
from listas import GestorListas
import listas

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
import os
import time

meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

def crearNuevoArchivo(i, a):
    """Crea un nuevo archivo Excel con todas las listas y la estructura base"""
    try:
        ruta = listas.crearFechaArchivo(i, a)
        
        # Cargar el archivo recién creado para agregar la hoja principal
        book = openpyxl.load_workbook(ruta)
        
        # Crear hoja principal
        hoja_principal = book.create_sheet('ATC - 001', 0)  # 0 para ponerla primera
        
        # Encabezados
        encabezados = ["FECHA", "CALLSING", "TYP", "REG", "Columna1", "GNSS", "RVSM", "CAT", "WTC", "DEP AD", "DEST AD"]
        for col, encabezado in enumerate(encabezados, 1):
            hoja_principal.cell(row=1, column=col, value=encabezado)
            hoja_principal.cell(row=1, column=col).font = Font(bold=True)
        
        book.save(ruta)
        return True, f"Archivo creado exitosamente: {ruta}"
    except Exception as e:
        return False, f"Error al crear archivo: {str(e)}"

def editarArchivoExistente(i, a, datos_vuelo):
    """Agrega un nuevo registro al archivo Excel existente"""
    try:
        ruta = os.path.join(os.path.expanduser('~'), 'Desktop', f'{meses[i]}_{a}.xlsx')
        
        if not os.path.exists(ruta):
            return False, "El archivo no existe. Crea uno nuevo primero."
        
        book = openpyxl.load_workbook(ruta)
        hoja = book['ATC - 001']
        
        # Encontrar la primera fila vacía
        fila_actual = 2
        while hoja[f'A{fila_actual}'].value is not None:
            fila_actual += 1
        
        # Agregar datos del vuelo
        hoja[f'A{fila_actual}'] = datos_vuelo['fecha']
        hoja[f'B{fila_actual}'] = datos_vuelo['callsign']
        hoja[f'C{fila_actual}'] = datos_vuelo['typ']
        hoja[f'D{fila_actual}'] = datos_vuelo['reg']
        
        # Fórmulas
        hoja[f'E{fila_actual}'] = f'=IFERROR(VLOOKUP(D{fila_actual},\'LISTA REG\'!A:B,2,FALSE),"")'
        hoja[f'F{fila_actual}'] = f'=IFERROR(VLOOKUP(C{fila_actual},\'GNSS - RVSM\'!A:C,2,FALSE),"")'
        hoja[f'G{fila_actual}'] = f'=IFERROR(VLOOKUP(C{fila_actual},\'GNSS - RVSM\'!A:C,3,FALSE),"")'
        hoja[f'H{fila_actual}'] = f'=IFERROR(VLOOKUP(C{fila_actual},\'LISTA CAT - WTC\'!A:C,2,FALSE),"")'
        hoja[f'I{fila_actual}'] = f'=IFERROR(VLOOKUP(C{fila_actual},\'LISTA CAT - WTC\'!A:C,3,FALSE),"")'
        
        hoja[f'J{fila_actual}'] = datos_vuelo['dep_ad']
        hoja[f'K{fila_actual}'] = datos_vuelo['dest_ad']
        
        book.save(ruta)
        return True, "Datos agregados exitosamente!"
        
    except Exception as e:
        return False, f"Error al editar archivo: {str(e)}"